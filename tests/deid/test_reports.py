"""Reports: reading, ID-only matching, redaction layers, header block (SPEC §6.2, §3.1; T19, T32).

TR-RPT-02, TR-RPT-03, TR-RPT-01, TR-RPT-04, TR-RPT-06.
"""

from __future__ import annotations

import socket
from collections import Counter
from pathlib import Path

import docx
import pytest
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

from app.deid.index import index_input
from app.deid.report import header, history, ner
from app.deid.report.match import hashed_name, match_reports
from app.deid.report.read import read_report
from app.deid.report.redact import redact_report, residual_patterns
from tests.deid.helpers import write_study


def test_sample_reports_matched(sample_inbox: Path) -> None:  # TR-RPT-01
    idx = index_input(sample_inbox)
    m = match_reports(idx)
    assert len(m.by_study) == 4 and m.unmatched == []
    assert sorted(m.tier.values()) == ["accession", "accession", "accession", "folder"]
    fmts = sorted((read_report(p).source_format, read_report(p).extraction) for p in m.by_study.values())
    assert fmts == [("docx", "native"), ("pdf", "native"), ("pdf", "ocr"), ("txt", "native")]


def test_docx_table_cells_read(tmp_path: Path) -> None:  # TR-RPT-02
    d = docx.Document()
    d.add_paragraph("CT THORAX")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Patient Name", "Mrs. Lakshmi Iyer"
    d.save(tmp_path / "r.docx")
    text = read_report(tmp_path / "r.docx").text
    assert "Lakshmi Iyer" in text and "CT THORAX" in text


def _image_page(lines: list[str]) -> Image.Image:
    img = Image.new("L", (1654, 2339), 255)
    draw = ImageDraw.Draw(img)
    for i, line in enumerate(lines):
        draw.text((120, 140 + 52 * i), line, fill=0, font=ImageFont.load_default(size=30))
    return img


def test_mixed_pdf_page_by_page(tmp_path: Path) -> None:  # TR-RPT-02
    path = tmp_path / "mixed.pdf"
    c = canvas.Canvas(str(path), pagesize=A4)
    c.setAuthor("Dr Neha Gupta")
    c.drawString(50, 780, "Page one has a real text layer with enough characters to count.")
    c.showPage()
    c.drawImage(ImageReader(_image_page(["SCANNED PAGE TWO", "Liver normal in size"])), 0, 0, *A4)
    c.save()
    rep = read_report(path)
    assert rep.extraction == "ocr" and rep.pages_ocr == 1
    assert "real text layer" in rep.text and "SCANNEDPAGETWO" in rep.text.upper().replace(" ", "")
    assert "Gupta" not in rep.text  # PDF metadata is never read


def test_titled_names_and_patterns() -> None:  # TR-RPT-03
    text = "Reported by Dr.Kavita Menon\nDR. ANIL DESHPANDE, MD\nCall +91 98220 45671, ananya.p@example.com"
    out, counts = redact_report(text, frozenset(), use_ner=False)
    for leak in ("Kavita", "Menon", "DESHPANDE", "98220", "example.com"):
        assert leak not in out
    assert counts["NAME"] == 2 and residual_patterns(out) == []


def test_header_identifiers_on_token_boundaries() -> None:  # TR-RPT-03
    out, counts = redact_report("Technique: plain. Tech Sunil present.", frozenset({"Tech", "Sunil"}), False)
    assert out.startswith("Technique") and "Sunil" not in out and counts["KNOWN_ID"] == 2


def test_t19_clinical_history() -> None:  # T19, TR-RPT-04
    text = "Clinical history: S/O Ramesh, farmer, Sardhana. Works at Modi Mills.\nFindings: Normal chest."
    counts: Counter[str] = Counter()
    out = history.redact(text, counts)
    for leak in ("Ramesh", "farmer", "Sardhana", "Modi"):
        assert leak not in out
    assert out.endswith("Findings: Normal chest.")  # findings unchanged
    assert counts["RELATIVE"] == 1 and counts["OCCUPATION"] == 1 and counts["PLACE"] == 1


def test_complaints_of_is_not_a_relative() -> None:  # TR-RPT-04
    out = history.redact("C/O Fever since 3 days.", Counter())
    assert out == "C/O Fever since 3 days."


def test_t32_name_only_report_unmatched(tmp_path: Path) -> None:  # T32, TR-RPT-06
    inbox = tmp_path / "in"
    common = {"name": "Kumar^Ravi", "PatientAge": "045Y", "ReferringPhysicianName": "Dr^Demo^Ref"}
    write_study(inbox / "a", patient_id="DEMO-UHID-1001", accession="ACC-T1001", **common)
    write_study(inbox / "b", patient_id="DEMO-UHID-1002", accession="ACC-T1002", **common)
    (inbox / "reports").mkdir()
    (inbox / "reports" / "Ravi_Kumar_45Y.txt").write_text("Name: Ravi Kumar  Age: 45Y\nRef: Dr Demo", "utf-8")
    m = match_reports(index_input(inbox))
    assert m.by_study == {} and m.unmatched == [hashed_name(inbox / "reports" / "Ravi_Kumar_45Y.txt")]


def test_accession_must_match_whole_token(tmp_path: Path) -> None:  # TR-RPT-06
    inbox = tmp_path / "in"
    write_study(inbox / "a", accession="ACC-1023")
    (inbox / "reports").mkdir()
    (inbox / "reports" / "ACC-10231_report.txt").write_text("x", "utf-8")
    assert match_reports(index_input(inbox)).by_study == {}


def test_ambiguous_reports_unmatched(tmp_path: Path) -> None:  # TR-RPT-06
    inbox = tmp_path / "in"
    write_study(inbox / "both", accession="ACC-T2001", patient_id="DEMO-P1", prefix="A")
    write_study(inbox / "both", accession="ACC-T2002", patient_id="DEMO-P2", prefix="B")
    (inbox / "both" / "report.txt").write_text("x", "utf-8")  # folder with 2 studies
    (inbox / "both" / "ACC-T2001_ACC-T2002.txt").write_text("x", "utf-8")  # names two studies
    write_study(inbox / "one", accession="ACC-T3001", patient_id="DEMO-P3")
    (inbox / "one" / "a.txt").write_text("x", "utf-8")  # two reports, same tier
    (inbox / "one" / "b.txt").write_text("x", "utf-8")
    m = match_reports(index_input(inbox))
    assert m.by_study == {} and len(m.unmatched) == 4


def test_patient_id_tier(tmp_path: Path) -> None:  # TR-RPT-01
    inbox = tmp_path / "in"
    write_study(inbox / "s1", accession="", patient_id="DEMO-UHID-5001")
    write_study(inbox / "s2", accession="", patient_id="DEMO-UHID-5002")
    write_study(inbox / "s3", accession="", patient_id="DEMO-UHID-5002")  # patient with two studies
    (inbox / "reports").mkdir()
    (inbox / "reports" / "DEMO-UHID-5001.txt").write_text("x", "utf-8")
    (inbox / "reports" / "DEMO-UHID-5002.txt").write_text("x", "utf-8")
    m = match_reports(index_input(inbox))
    assert list(m.tier.values()) == ["patient_id"] and len(m.unmatched) == 1


def test_header_round_trip() -> None:  # TR-RPT-02
    text = header.build_report("S0123456789AB", "P0123456789AB", ["S0123456789AB-0002-000001"], "Body\n")
    parsed = header.parse_header(text)
    assert parsed == header.ReportHeader("S0123456789AB", "P0123456789AB", ("S0123456789AB-0002-000001",))
    assert text.endswith("Body\n") and header.parse_header("garbage") is None


def test_ner_makes_no_network_calls(monkeypatch: pytest.MonkeyPatch) -> None:  # TR-RPT-03, CLAUDE.md rule 15
    def refuse(*_a: object, **_k: object) -> None:
        raise AssertionError("outbound network call")

    monkeypatch.setattr(socket.socket, "connect", refuse)
    monkeypatch.setattr(socket, "create_connection", refuse)
    monkeypatch.setattr(ner, "_engine", None)
    counts: Counter[str] = Counter()
    out = ner.redact("The patient Priya Sharma was seen in Meerut. Fleischner follow-up advised.", counts)
    assert "Priya" not in out and "Fleischner" in out and counts["NER_PERSON"] >= 1
