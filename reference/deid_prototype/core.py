"""
Core de-identification rules shared by the engine: tag allowlist, keyed pseudonyms, date shift,
report redaction patterns. (Originally the PoC batch pipeline; the batch runner now lives in engine.py.)
"""
import argparse, csv, datetime as dt, hashlib, hmac, json, os, random, re, secrets, shutil, sys
from collections import defaultdict
from pathlib import Path

import pydicom
from pydicom.dataset import Dataset, FileMetaDataset
from pydicom.uid import generate_uid, PYDICOM_IMPLEMENTATION_UID

HERE = Path(__file__).resolve().parent

# --------------------------------------------------------------------------------------------------
# Tag policy
# --------------------------------------------------------------------------------------------------
KEEP_TAGS = {
    # image pixel / display
    "SamplesPerPixel", "PhotometricInterpretation", "Rows", "Columns", "BitsAllocated", "BitsStored",
    "HighBit", "PixelRepresentation", "PlanarConfiguration", "NumberOfFrames", "PixelData",
    "RescaleIntercept", "RescaleSlope", "RescaleType", "WindowCenter", "WindowWidth", "VOILUTFunction",
    "PixelSpacing", "ImagerPixelSpacing", "PixelAspectRatio", "LossyImageCompression",
    "LossyImageCompressionRatio", "LossyImageCompressionMethod", "PresentationLUTShape",
    # geometry
    "ImagePositionPatient", "ImageOrientationPatient", "SliceThickness", "SpacingBetweenSlices",
    "SliceLocation", "PatientPosition", "ViewPosition", "Laterality", "ImageLaterality",
    # acquisition / equipment (useful for AI buyers, not identifying)
    "Modality", "BodyPartExamined", "KVP", "ExposureTime", "XRayTubeCurrent", "Exposure",
    "ConvolutionKernel", "ContrastBolusAgent", "Manufacturer", "ManufacturerModelName",
    "ImageType", "SOPClassUID", "SeriesNumber", "InstanceNumber", "AcquisitionNumber",
    "StudyDescription", "SeriesDescription", "ProtocolName", "BurnedInAnnotation",
    # demographics kept at low granularity
    "PatientSex", "PatientAge",
}
DATE_TAGS = ["StudyDate", "SeriesDate", "AcquisitionDate", "ContentDate"]
UID_TAGS = ["StudyInstanceUID", "SeriesInstanceUID", "SOPInstanceUID", "FrameOfReferenceUID"]
# free-text tags we keep but still scan for leaked identifiers
SCAN_TEXT_TAGS = ["StudyDescription", "SeriesDescription", "ProtocolName", "ContrastBolusAgent"]

# --------------------------------------------------------------------------------------------------
# Report redaction rules (India-oriented)
# --------------------------------------------------------------------------------------------------
LABEL_LINE = re.compile(
    r"(?im)^(\s*(?:patient\s*name|pt\.?\s*name|name|uhid|mrn|reg(?:istration)?\.?\s*no\.?|patient\s*id|"
    r"ip\s*no\.?|op\s*no\.?|lab\s*no\.?|accession(?:\s*no\.?)?|ref(?:erred)?\.?\s*by|referring\s*(?:doctor|physician)|"
    r"address|mobile|phone|contact(?:\s*no\.?)?|email|aadhaar|abha(?:\s*(?:id|no\.?))?|d\.?o\.?b\.?|date\s*of\s*birth)"
    r"\s*[:\-]\s*)(.+)$")
REGEX_RULES = [
    ("EMAIL",   re.compile(r"\b[\w.+-]+@[\w-]+\.[\w.-]+\b")),
    ("PHONE",   re.compile(r"(?<!\d)(?:\+?91[\s-]?)?[6-9]\d{4}[\s-]?\d{5}(?!\d)")),
    ("ID12",    re.compile(r"(?<!\d)\d{4}[\s-]?\d{4}[\s-]?\d{4}(?!\d)")),       # Aadhaar-like
    ("ABHA",    re.compile(r"(?<!\d)\d{2}-\d{4}-\d{4}-\d{4}(?!\d)")),
    ("PAN",     re.compile(r"\b[A-Z]{5}\d{4}[A-Z]\b")),
    ("REGNO",   re.compile(r"\b(?:Reg(?:istration)?|IP|OP|Lab)\.?\s*(?:No\.?|Number)\s*[:#-]?\s*(?=[A-Z0-9/\-]*\d)[A-Z0-9][A-Z0-9/\-]{3,}", re.I)),
    ("MRN",     re.compile(r"\b(?:UHID|MRN|ABHA)\s*(?:No\.?|ID)?\s*[:#-]?\s*(?=\S*\d)[A-Z0-9][A-Z0-9/\-]{3,}", re.I)),
    ("DATE",    re.compile(r"\b\d{1,2}[/.-]\d{1,2}[/.-]\d{2,4}\b|\b\d{1,2}\s+(?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\.?,?\s+\d{2,4}\b", re.I)),
    ("PIN",     re.compile(r"(?<!\d)[1-9]\d{5}(?!\d)")),
    ("NAME",    re.compile(r"\b(?:Dr|Mr|Mrs|Ms|Miss|Smt|Shri|Sri|Kumari|Master|Baby)(?:\.\s*|\s+)[A-Z][a-zA-Z]+(?:\s+[A-Z]\.)?(?:\s+[A-Z][a-zA-Z]+){0,2}")),
]

# --------------------------------------------------------------------------------------------------
# Helpers
# --------------------------------------------------------------------------------------------------
def load_config(path):
    with open(path) as f:
        cfg = json.load(f)
    for k in ("input_dir", "output_dir", "secret_key_file"):
        p = Path(cfg[k])
        cfg[k] = p if p.is_absolute() else (Path(path).parent / p)
    return cfg

def load_or_create_key(path: Path) -> bytes:
    if path.exists():
        return bytes.fromhex(path.read_text().strip())
    path.parent.mkdir(parents=True, exist_ok=True)
    key = secrets.token_bytes(32)
    path.write_text(key.hex())
    try:
        os.chmod(path, 0o600)
    except OSError:
        pass
    print(f"[key] new secret key created at {path} - back it up securely; never ship it with data")
    return key

def h(key: bytes, kind: str, value: str) -> str:
    return hmac.new(key, f"{kind}:{value}".encode(), hashlib.sha256).hexdigest()

def pseudo_uid(key, original):
    # 2.25.<decimal of 120-bit int> -> valid DICOM UID, < 64 chars
    return "2.25." + str(int(h(key, "uid", original)[:30], 16))

def sha256_file(p: Path) -> str:
    d = hashlib.sha256()
    with open(p, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            d.update(chunk)
    return d.hexdigest()

def parse_da(s):
    try:
        return dt.datetime.strptime(str(s)[:8], "%Y%m%d").date()
    except Exception:
        return None

def age_from(ds):
    age = str(ds.get("PatientAge", "") or "")
    if re.fullmatch(r"\d{3}[DWMY]", age):
        return age
    dob, sd = parse_da(ds.get("PatientBirthDate", "")), parse_da(ds.get("StudyDate", ""))
    if dob and sd:
        yrs = sd.year - dob.year - ((sd.month, sd.day) < (dob.month, dob.day))
        return f"{max(yrs, 0):03d}Y"
    return ""

def _norm(v):
    """Lower-case, punctuation-insensitive form used to compare vendor names."""
    return re.sub(r"[^a-z0-9]+", " ", str(v or "").lower()).strip()


def known_identifiers(ds):
    """Exact strings from the original header that must never appear in outputs."""
    vals = set()
    # Some modality software (e.g. dental CBCT) writes the vendor's own name into InstitutionName.
    # The vendor name is not a patient or site identifier and is kept in Manufacturer for buyers,
    # so an institution value that is really the vendor name is not treated as an identifier.
    vendor = {_norm(ds.get(k, "")) for k in ("Manufacturer", "ManufacturerModelName")} - {""}
    for tag in ("PatientID", "AccessionNumber", "OtherPatientIDs", "InstitutionName",
                "InstitutionAddress", "PatientAddress", "PatientTelephoneNumbers"):
        v = str(ds.get(tag, "") or "").strip()
        if tag.startswith("Institution") and _norm(v) in vendor:
            continue
        if len(v) >= 3:
            vals.add(v)
    for tag in ("PatientName", "ReferringPhysicianName", "PerformingPhysicianName",
                "OperatorsName", "NameOfPhysiciansReadingStudy"):
        v = ds.get(tag, None)
        if v is None:
            continue
        for part in re.split(r"[\^\s,]+", str(v)):
            if len(part) >= 3 and not part.lower() in {"dr", "mr", "mrs", "smt", "shri"}:
                vals.add(part)
    dob = parse_da(ds.get("PatientBirthDate", ""))
    if dob:
        vals.update({dob.strftime("%d/%m/%Y"), dob.strftime("%d-%m-%Y"), dob.strftime("%Y%m%d"),
                     dob.strftime("%d.%m.%Y")})
    return vals

def is_quarantine(ds, cfg):
    mod = str(ds.get("Modality", "")).upper()
    if mod in cfg["quarantine_modalities"]:
        return f"modality {mod}"
    if str(ds.get("BurnedInAnnotation", "")).upper() == "YES":
        return "BurnedInAnnotation=YES"
    sop = str(ds.get("SOPClassUID", ""))
    if sop.startswith("1.2.840.10008.5.1.4.1.1.7"):          # secondary capture family
        return "secondary capture (e.g. CT dose screen)"
    if sop == "1.2.840.10008.5.1.4.1.1.104.1":               # encapsulated PDF
        return "encapsulated PDF"
    itype = [str(x).upper() for x in (ds.get("ImageType", []) or [])]
    if "SECONDARY" in itype and "SCREEN SAVE" in " ".join(itype):
        return "screen save"
    return None

# --------------------------------------------------------------------------------------------------
# DICOM de-identification
# --------------------------------------------------------------------------------------------------
def deid_dataset(ds, key, cfg, pat_pseudo, study_pseudo, shift_days):
    out = Dataset()
    for kw in KEEP_TAGS:
        if kw in ds:
            out[kw] = ds[kw]
    # scrub identifiers typed into free-text descriptions (e.g. "CHEST PA - RAMESH")
    idents = sorted((i for i in known_identifiers(ds) if len(i) >= 3), key=len, reverse=True)
    for kw in SCAN_TEXT_TAGS:
        if kw in out and idents:
            val = str(out[kw].value)
            for i in idents:
                val = re.sub(re.escape(i), "[REDACTED]", val, flags=re.I)
            el = out[kw]
            out[kw] = pydicom.DataElement(el.tag, el.VR, val[:64])   # new element; source stays untouched
    # viewer compatibility: a zero/invalid pixel spacing (seen in real exports) makes web viewers such as
    # OHIF render a black image, so drop it rather than pass it on
    for kw in ("PixelSpacing", "ImagerPixelSpacing"):
        if kw in out:
            try:
                ok = all(float(v) > 0 for v in out[kw].value)
            except Exception:
                ok = False
            if not ok:
                del out[kw]
    # identity
    out.PatientID = pat_pseudo
    out.PatientName = pat_pseudo
    out.AccessionNumber = study_pseudo
    out.StudyID = study_pseudo[-8:]
    age = age_from(ds)
    if age:
        if age.endswith("Y") and int(age[:3]) >= cfg["age_cap_years"]:
            age = f"{cfg['age_cap_years']:03d}Y"
        out.PatientAge = age
    # UIDs
    for kw in UID_TAGS:
        if kw in ds:
            setattr(out, kw, pseudo_uid(key, str(ds.get(kw))))
    # dates
    for kw in DATE_TAGS:
        d = parse_da(ds.get(kw, ""))
        if not d or cfg["date_mode"] == "remove":
            continue
        if cfg["date_mode"] == "year_only":
            setattr(out, kw, f"{d.year}0101")
        else:
            setattr(out, kw, (d - dt.timedelta(days=shift_days)).strftime("%Y%m%d"))
    # de-identification flags (DICOM PS3.15)
    out.PatientIdentityRemoved = "YES"
    out.DeidentificationMethod = "PoC allowlist; HMAC pseudonyms; UIDs remapped; dates " + cfg["date_mode"]
    # file meta
    meta = FileMetaDataset()
    meta.MediaStorageSOPClassUID = ds.get("SOPClassUID", "1.2.840.10008.5.1.4.1.1.7")
    meta.MediaStorageSOPInstanceUID = out.get("SOPInstanceUID", generate_uid())
    meta.TransferSyntaxUID = ds.file_meta.TransferSyntaxUID
    meta.ImplementationClassUID = PYDICOM_IMPLEMENTATION_UID
    out.file_meta = meta
    return out

# --------------------------------------------------------------------------------------------------
# Report de-identification
# --------------------------------------------------------------------------------------------------
def read_report(p: Path) -> str:
    """Plain text of a report. For PDFs see read_report_ex (tells you whether OCR was needed)."""
    return read_report_ex(p)[0]


def read_report_ex(p: Path, min_chars_per_page=40, ocr_dpi=300):
    """Returns (text, source_format, extraction) with extraction in {"native", "ocr"}.

    .txt  -> read as UTF-8
    .docx -> paragraphs + table cells (many Indian report templates put patient details in a table)
    .pdf  -> text layer via pypdfium2; pages with too little text (scanned reports) are rendered
             at ocr_dpi and OCR'd with RapidOCR, lines rebuilt top-to-bottom. PDF metadata
             (author, title, producer) is ignored and the original PDF is never copied to output.
    """
    suf = p.suffix.lower()
    if suf == ".docx":
        import docx
        d = docx.Document(str(p))
        parts = [par.text for par in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append("  ".join(c.text.strip() for c in row.cells))
        return "\n".join(parts), "docx", "native"
    if suf == ".pdf":
        import pypdfium2 as pdfium
        pdf = pdfium.PdfDocument(str(p))
        pages, used_ocr = [], False
        for page in pdf:
            text = page.get_textpage().get_text_range().replace("\r\n", "\n")
            if len(text.strip()) < min_chars_per_page:
                text = _ocr_page(page, ocr_dpi)
                used_ocr = True
            pages.append(text)
        return "\n\f\n".join(pages), "pdf", "ocr" if used_ocr else "native"
    return p.read_text(encoding="utf-8", errors="replace"), "txt", "native"


def _ocr_page(page, dpi):
    import numpy as np
    from rapidocr_onnxruntime import RapidOCR
    global _PDF_OCR
    try:
        _PDF_OCR
    except NameError:
        _PDF_OCR = RapidOCR()
    img = page.render(scale=dpi / 72).to_pil().convert("RGB")
    res, _ = _PDF_OCR(np.array(img))
    items = sorted(((min(p[1] for p in b), min(p[0] for p in b), t) for b, t, _s in (res or [])))
    tol = 12 * dpi / 200
    lines, cur, start_y = [], [], None
    for y, x, t in items:                      # group boxes into lines by vertical position
        if start_y is not None and y - start_y > tol:
            lines.append(" ".join(w for _, w in sorted(cur)))
            cur, start_y = [], None
        if start_y is None:
            start_y = y
        cur.append((x, t))
    if cur:
        lines.append(" ".join(w for _, w in sorted(cur)))
    return "\n".join(lines)

def redact_report(text, identifiers):
    counts = defaultdict(int)
    def lab(m):
        counts["LABELLED_FIELD"] += 1
        return m.group(1) + "[REDACTED]"
    text = LABEL_LINE.sub(lab, text)
    for ident in sorted(identifiers, key=len, reverse=True):
        pat = re.compile(re.escape(ident), re.I)
        text, n = pat.subn("[REDACTED]", text)
        counts["KNOWN_ID"] += n
    for name, rx in REGEX_RULES:
        text, n = rx.subn(f"[{name}]", text)
        counts[name] += n
    return text, dict(counts)

def find_reports(input_dir, exts):
    return [p for p in input_dir.rglob("*") if p.is_file() and p.suffix.lower() in exts]

# --------------------------------------------------------------------------------------------------
# Main run
# --------------------------------------------------------------------------------------------------
