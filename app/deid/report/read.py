"""Read a report as plain text (SPEC §6.2 "Reading"). TR-RPT-01, TR-RPT-02.

Port of ``reference/deid_prototype/core.read_report_ex``:
- .txt  read as UTF-8;
- .docx paragraphs and table cells (many Indian templates put patient details in a table). Section headers
  and footers (letterheads) are not read, so they can never reach the output;
- .pdf  text layer via pypdfium2, page by page; a page with fewer than 40 characters (a scanned page) is
  rendered at 300 dpi and OCR'd with RapidOCR, lines rebuilt top to bottom.
PDF/DOCX metadata (author, title, producer) is never read, and the source file is never copied.
"""

from __future__ import annotations

from pathlib import Path

import numpy as np

from app.deid.ocr import engine
from app.deid.types import DeidError, ExtractedReport

MIN_CHARS_PER_PAGE = 40
OCR_DPI = 300


def read_report(path: Path) -> ExtractedReport:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return ExtractedReport(_read_docx(path), "docx", "native")
    if suffix == ".pdf":
        return _read_pdf(path)
    if suffix == ".txt":
        return ExtractedReport(path.read_text(encoding="utf-8", errors="replace"), "txt", "native")
    raise DeidError("unsupported report format (convert .doc/.rtf first)")


def _read_docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append("  ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _read_pdf(path: Path) -> ExtractedReport:
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(str(path))
    pages, ocr_pages = [], 0
    try:
        for page in pdf:
            text = page.get_textpage().get_text_range().replace("\r\n", "\n")
            if len(text.strip()) < MIN_CHARS_PER_PAGE:
                text = _ocr_page(page)
                ocr_pages += 1
            pages.append(text)
    finally:
        pdf.close()
    return ExtractedReport("\n\f\n".join(pages), "pdf", "ocr" if ocr_pages else "native", ocr_pages)


def _ocr_page(page: object) -> str:
    image = page.render(scale=OCR_DPI / 72).to_pil().convert("RGB")  # type: ignore[attr-defined]
    items = sorted(
        (min(p[1] for p in box), min(p[0] for p in box), text)
        for box, text, _score in engine.run(np.array(image))
    )
    return "\n".join(group_lines(items, tolerance=12 * OCR_DPI / 200))


def group_lines(items: list[tuple[float, float, str]], tolerance: float) -> list[str]:
    """Group (y, x, text) boxes into lines by vertical position, words left to right."""
    lines: list[str] = []
    current: list[tuple[float, str]] = []
    start_y: float | None = None
    for y, x, text in items:
        if start_y is not None and y - start_y > tolerance:
            lines.append(" ".join(w for _, w in sorted(current)))
            current, start_y = [], None
        if start_y is None:
            start_y = y
        current.append((x, text))
    if current:
        lines.append(" ".join(w for _, w in sorted(current)))
    return lines
